# pyrefly: ignore [missing-import]
from celery import shared_task
# pyrefly: ignore [missing-import]
from django.core.mail import EmailMessage
# pyrefly: ignore [missing-import]
from django.contrib.auth import get_user_model
# pyrefly: ignore [missing-import]
from apps.transactions.models import Transaction
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import io
# pyrefly: ignore [missing-import]
from django.db.models import Sum

User = get_user_model()


@shared_task(bind=True)
def generate_pdf_report(self, user_id, month, year):
    """Generate PDF expense report and email it to the user"""
    user = User.objects.get(id=user_id)
    transactions = Transaction.objects.filter(
        user=user, date__month=month, date__year=year
    ).select_related('category').order_by('date')

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []

    # Title
    story.append(Paragraph(f"Expense Report — {month}/{year}", styles['Title']))
    story.append(Paragraph(f"Account: {user.email}", styles['Normal']))
    story.append(Spacer(1, 16))

    # Summary
    total_income = transactions.filter(type='income').aggregate(s=Sum('amount'))['s'] or 0
    total_expense = transactions.filter(type='expense').aggregate(s=Sum('amount'))['s'] or 0
    balance = total_income - total_expense

    summary_data = [
        ['Metric', 'Amount'],
        ['Total Income', f"{user.currency} {total_income:.2f}"],
        ['Total Expense', f"{user.currency} {total_expense:.2f}"],
        ['Net Balance', f"{user.currency} {balance:.2f}"],
    ]
    summary_table = Table(summary_data, colWidths=[200, 200])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor('#EEF2FF'), colors.white]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 20))

    # Transactions table
    data = [['Date', 'Title', 'Category', 'Type', 'Amount']]
    for tx in transactions:
        data.append([
            str(tx.date),
            tx.title[:30],
            tx.category.name if tx.category else '—',
            tx.type.capitalize(),
            f"{user.currency} {tx.amount:.2f}",
        ])

    table = Table(data, colWidths=[70, 150, 100, 70, 100])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F3F4F6')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
    ]))
    story.append(table)

    doc.build(story)
    pdf_data = buffer.getvalue()
    buffer.close()

    email = EmailMessage(
        subject=f"Your Expense Report — {month}/{year}",
        body=(
            f"Hi {user.username},\n\n"
            f"Please find your expense report for {month}/{year} attached.\n\n"
            f"Summary:\n"
            f"  Income:  {user.currency} {total_income:.2f}\n"
            f"  Expense: {user.currency} {total_expense:.2f}\n"
            f"  Balance: {user.currency} {balance:.2f}\n\n"
            f"— Expense Tracker"
        ),
        to=[user.email],
    )
    email.attach(f"expense_report_{month}_{year}.pdf", pdf_data, 'application/pdf')
    email.send()
    return f"PDF report sent to {user.email}"


@shared_task(bind=True)
def generate_excel_report(self, user_id, month, year):
    """Generate Excel expense report and email it"""
    user = User.objects.get(id=user_id)
    transactions = Transaction.objects.filter(
        user=user, date__month=month, date__year=year
    ).select_related('category').order_by('date')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Report {month} {year}"

    # Header style
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')

    headers = ['Date', 'Title', 'Category', 'Type', 'Amount', 'Currency', 'Note']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for row, tx in enumerate(transactions, 2):
        ws.cell(row=row, column=1, value=str(tx.date))
        ws.cell(row=row, column=2, value=tx.title)
        ws.cell(row=row, column=3, value=tx.category.name if tx.category else '—')
        ws.cell(row=row, column=4, value=tx.type.capitalize())
        ws.cell(row=row, column=5, value=float(tx.amount))
        ws.cell(row=row, column=6, value=user.currency)
        ws.cell(row=row, column=7, value=tx.note)

    # Auto-width columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    excel_data = buffer.getvalue()
    buffer.close()

    email = EmailMessage(
        subject=f"Your Excel Report — {month}/{year}",
        body=f"Hi {user.username},\n\nYour Excel report for {month}/{year} is attached.\n\n— Expense Tracker",
        to=[user.email],
    )
    email.attach(
        f"expense_report_{month}_{year}.xlsx",
        excel_data,
        'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    email.send()
    return f"Excel report sent to {user.email}"
